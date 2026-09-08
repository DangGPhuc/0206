"""
0206 - Report Template Validator Module
Phase 21: Validates user-provided DOCX report templates for structural and schema integrity.
"""
from pathlib import Path
from typing import Tuple, List
import docx


class TemplateValidationError(Exception):
    """Raised when a report template fails validation."""
    pass


class TemplateValidator:
    """Validates DOCX templates for compatibility with report adapters."""

    @staticmethod
    def validate_sans_style_template(template_path: Path) -> Tuple[bool, List[str]]:
        """
        Checks if a template matches the SANS-style 58-row table structure.
        """
        p = Path(template_path)
        warnings = []

        if not p.exists():
            return False, [f"Template file does not exist: {p}"]

        if p.suffix.lower() != ".docx":
            return False, [f"Template must be a .docx document, got: {p.suffix}"]

        try:
            doc = docx.Document(str(p))
        except Exception as e:
            return False, [f"Failed to open DOCX file: {e}"]

        if not doc.tables:
            return False, ["Template contains no tables. Expected at least 1 table."]

        table = doc.tables[0]
        row_count = len(table.rows)

        if row_count < 40:
            return False, [f"Template table has only {row_count} rows. SANS-style requires at least 40-58 rows."]

        if row_count != 58:
            warnings.append(f"Template table has {row_count} rows (standard template has 58 rows). Adapters will map available rows.")

        return True, warnings
