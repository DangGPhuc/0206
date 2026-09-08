"""
0206 - SANS-Style DOCX Report Adapter
Phase 20 & 21: Populates findings into user-provided SANS FOR610-style DOCX templates.
Distinguishes evidence states: [OBSERVED], [INFERRED], [NOT_CONFIRMED], [NOT_ANALYZED].
"""
from pathlib import Path
from typing import Dict, Any, Optional
import docx
from docx.shared import Pt, RGBColor
from reporting.adapters.base import BaseReportAdapter
from reporting.validators.template_validator import TemplateValidator, TemplateValidationError


class SANSStyleReportAdapter(BaseReportAdapter):
    """Fills evidence-grounded findings into a 58-row SANS-style template."""

    def __init__(self, template_path: Path):
        self.template_path = Path(template_path)
        valid, warnings = TemplateValidator.validate_sans_style_template(self.template_path)
        if not valid:
            raise TemplateValidationError("; ".join(warnings))

    def _set_cell_text(self, cell, text: str, font_name: str = "Calibri", font_size_pt: float = 9.5, bold: bool = False):
        cell.text = ""
        lines = text.strip().split("\n")
        for i, line in enumerate(lines):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(line)
            run.font.name = font_name
            run.font.size = Pt(font_size_pt)
            run.bold = bold

    def render(self, session_data: Dict[str, Any], output_path: Path) -> Path:
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        doc = docx.Document(str(self.template_path))
        table = doc.tables[0]
        rows = table.rows

        manifest = session_data.get("manifest", {})
        assessment = session_data.get("assessment", {})
        findings = session_data.get("findings", [])

        def set_kv(r_idx: int, c_idx: int, val: str):
            if r_idx < len(rows) and c_idx < len(rows[r_idx].cells):
                self._set_cell_text(rows[r_idx].cells[c_idx], val)

        def set_data(r_idx: int, val: str):
            if r_idx < len(rows):
                c = rows[r_idx].cells[1] if len(rows[r_idx].cells) > 1 else rows[r_idx].cells[0]
                self._set_cell_text(c, val)

        # Overview Metadata
        set_data(1, manifest.get("sample_filename", "N/A"))
        hashes = manifest.get("sample_hashes", {})
        set_data(2, hashes.get("sha256", "N/A"))
        set_data(3, hashes.get("sha1", "N/A"))
        set_data(4, hashes.get("md5", "N/A"))
        set_data(5, f"{(manifest.get('sample_size_bytes') or 0):,} bytes")
        set_data(6, f"{assessment.get('threat_level', 'UNKNOWN')} ({assessment.get('threat_score', 0)}/100)")
        set_data(7, assessment.get("classification", "Generic"))

        # Executive narrative
        set_data(9, assessment.get("summary", "No summary generated."))

        # Findings summary
        findings_str = "\n".join([
            f"[{f.get('status') or f.get('evidence_level', 'CAPABILITY')}] {f.get('title')} (Evidence: {', '.join(f.get('source_evidence_ids', []) or f.get('evidence_ids', [])) or 'None'})"
            for f in findings[:15]
        ]) or "[NOT_ANALYZED] No findings registered."
        set_data(12, findings_str)

        # Grounded IOCs
        host_str = "\n".join(assessment.get("host_iocs", [])) or "None identified."
        net_str = "\n".join(assessment.get("network_iocs", [])) or "None identified."
        set_data(14, host_str)
        set_data(15, net_str)

        doc.save(str(output_path))
        return output_path
