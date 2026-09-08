"""
SANS FOR610 Malware Analysis Report DOCX Generator
Fills findings into Malware_Analysis_Report_Template.docx with structured precision.
"""
import shutil
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, Any, Optional
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from config import DEFAULT_TEMPLATE_PATH


class FOR610ReportGenerator:
    """Generates professional SANS FOR610 DOCX malware reports from analysis telemetry."""

    def __init__(self, template_path: Optional[Path] = None):
        self.template_path = Path(template_path) if template_path else DEFAULT_TEMPLATE_PATH
        if not self.template_path.exists():
            raise FileNotFoundError(f"Template file not found: {self.template_path}")

    def _set_cell_text(self, cell, text: str, font_name: str = "Calibri", font_size_pt: float = 9.5, bold: bool = False, color_rgb: Optional[RGBColor] = None):
        """Sets cell text with clean styling and paragraph formatting."""
        cell.text = ""
        lines = text.strip().split("\n")
        for i, line in enumerate(lines):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(line)
            run.font.name = font_name
            run.font.size = Pt(font_size_pt)
            run.bold = bold
            if color_rgb:
                run.font.color.rgb = color_rgb

    def generate(
        self,
        static_data: Dict[str, Any],
        behavioral_data: Dict[str, Any],
        ai_data: Dict[str, Any],
        output_path: Path
    ) -> Path:
        """Populates the SANS FOR610 DOCX template with combined findings."""
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        doc = docx.Document(str(self.template_path))
        if not doc.tables:
            raise ValueError("Template does not contain required tables.")

        table = doc.tables[0]
        rows = table.rows
        file_info = static_data.get("file_info", {})
        network = behavioral_data.get("network", {})
        host_beh = behavioral_data.get("host_behavior", {})

        def set_kv_row(row_idx: int, val_cell_idx: int, text: str):
            """Sets key-value row text."""
            if row_idx < len(rows) and val_cell_idx < len(rows[row_idx].cells):
                self._set_cell_text(rows[row_idx].cells[val_cell_idx], text)

        def set_data_row(row_idx: int, text: str):
            """Sets single merged data row text."""
            if row_idx < len(rows):
                self._set_cell_text(rows[row_idx].cells[0], text)

        # ----------------------------------------------------
        # 1. BACKGROUND (Rows 1 - 6)
        # ----------------------------------------------------
        current_date_str = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC")
        set_kv_row(1, 2, current_date_str)
        set_kv_row(2, 2, "RE-ANALYSIS-STATION-01 (Isolated Sandbox Env)")
        set_kv_row(3, 2, file_info.get("file_name", "Unknown"))
        set_kv_row(4, 2, file_info.get("file_path", "N/A"))
        
        timestamp_str = f"Compile: {file_info.get('compile_time', 'N/A')}"
        set_kv_row(5, 2, timestamp_str)
        set_kv_row(6, 2, ai_data.get("root_cause", "Suspicious download / Email attachment"))

        # ----------------------------------------------------
        # 2. STATIC ANALYSIS (Rows 8 - 28)
        # ----------------------------------------------------
        set_kv_row(8, 1, f"{file_info.get('file_size', 0):,} bytes ({file_info.get('file_size', 0)/1024:.2f} KB)")
        set_kv_row(9, 1, "Default application icon (PE embedded resource)")
        set_kv_row(10, 1, "Signed (Embedded Authenticode)" if file_info.get("is_signed") else "Unsigned (No digital signature found)")
        
        hashes_text = (
            f"MD5:    {file_info.get('md5', 'N/A')}\n"
            f"SHA1:   {file_info.get('sha1', 'N/A')}\n"
            f"SHA256: {file_info.get('sha256', 'N/A')}"
        )
        set_kv_row(11, 1, hashes_text)
        set_kv_row(12, 1, file_info.get("imphash", "N/A"))

        # Row 14: PE Section Hashes & Details
        sections = static_data.get("sections", [])
        if sections:
            sec_lines = [
                f"{'Name':<10} {'VirtSize':<10} {'RawSize':<10} {'Entropy':<9} {'Perms':<6} {'MD5 Hash':<32}"
            ]
            sec_lines.append("-" * 80)
            for s in sections:
                rwx_flag = " [RWX!]" if s.get("is_rwx") else ""
                sec_lines.append(
                    f"{s.get('name', ''):<10} {s.get('virtual_size', 0):<10} {s.get('raw_size', 0):<10} "
                    f"{s.get('entropy', 0.0):<9.2f} {s.get('permissions', ''):<6} {s.get('md5', ''):<32}{rwx_flag}"
                )
            set_data_row(14, "\n".join(sec_lines))
        else:
            set_data_row(14, "No PE sections identified.")

        # Row 16: Compile Time
        set_data_row(16, f"TimeDateStamp: {file_info.get('compile_time', 'N/A')}")

        # Row 18: File Properties
        props_text = (
            f"Architecture:    {file_info.get('architecture', 'Unknown')}\n"
            f"Subsystem:       {file_info.get('subsystem', 'Unknown')}\n"
            f"Entry Point:     {file_info.get('entry_point', '0x0')}\n"
            f"Image Base:      {file_info.get('image_base', '0x0')}\n"
            f"Total Sections:  {len(sections)}"
        )
        set_data_row(18, props_text)

        # Row 20: Strings
        strings_data = static_data.get("strings_analysis", {})
        strings_lines = []
        if strings_data.get("urls"):
            strings_lines.append(f"URLs ({len(strings_data['urls'])}):\n  " + "\n  ".join(strings_data["urls"][:8]))
        if strings_data.get("ips"):
            strings_lines.append(f"IPs ({len(strings_data['ips'])}):\n  " + "\n  ".join(strings_data["ips"][:8]))
        if strings_data.get("registry_keys"):
            strings_lines.append(f"Registry Keys ({len(strings_data['registry_keys'])}):\n  " + "\n  ".join(strings_data["registry_keys"][:6]))
        if strings_data.get("suspicious_commands"):
            strings_lines.append(f"Suspicious Commands/Strings ({len(strings_data['suspicious_commands'])}):\n  " + "\n  ".join(strings_data["suspicious_commands"][:8]))
        
        if not strings_lines:
            strings_lines.append(f"Extracted {strings_data.get('total_strings_count', 0)} total strings; no explicit plain-text URLs or registry keys matched.")
        set_data_row(20, "\n\n".join(strings_lines))

        # Row 22: Packed
        packed_info = static_data.get("packed_analysis", {})
        if packed_info.get("is_packed"):
            packed_text = (
                f"STATUS: PACKED / OBFUSCATED\n"
                f"Indicators:\n  - " + "\n  - ".join(packed_info.get("reasons", ["High entropy"]))
            )
        else:
            packed_text = "STATUS: NOT PACKED (Standard PE section layout, typical compiler entropy distribution)"
        set_data_row(22, packed_text)

        # Row 24: Entropy
        entropy_lines = [f"Overall File Entropy: {file_info.get('overall_entropy', 0.0):.4f} (Threshold: 7.0000)"]
        for s in sections:
            entropy_lines.append(f"  Section {s.get('name', '')}: {s.get('entropy', 0.0):.4f} {'(HIGH)' if s.get('entropy', 0) >= 7.0 else '(Normal)'}")
        set_data_row(24, "\n".join(entropy_lines))

        # Row 26: Imported / Exported Functions & API Hashing
        imp_lines = []
        caps = static_data.get("detected_capabilities", {})
        if caps:
            imp_lines.append("SUSPICIOUS CAPABILITIES FROM IMPORTED APIs:")
            for cap, apis in caps.items():
                imp_lines.append(f"  [{cap}]: {', '.join(apis)}")
            imp_lines.append("")

        imports = static_data.get("imports", {})
        if imports:
            imp_lines.append(f"IMPORTED DLLs ({len(imports)}):")
            for dll, funcs in list(imports.items())[:8]:
                imp_lines.append(f"  {dll.upper()} ({len(funcs)} functions): {', '.join(funcs[:6])}{'...' if len(funcs)>6 else ''}")
        else:
            imp_lines.append("No imports listed in Import Address Table.")

        # Add API Hashing constant resolutions if found
        api_hashes = static_data.get("api_hashing_matches", [])
        if api_hashes:
            imp_lines.append(f"\nDETECTED API HASHING CONSTANTS ({len(api_hashes)} resolved):")
            for m in api_hashes[:10]:
                imp_lines.append(f"  Hash {m['hash_hex']} -> {m['api']} (Algorithm: {m['algorithm']} at offset {m['offset_hex']})")

        set_data_row(26, "\n".join(imp_lines))

        # Row 28: Open Source Research
        osint_text = (
            f"VirusTotal Search Query: https://www.virustotal.com/gui/file/{file_info.get('sha256', '')}\n"
            f"MalwareBazaar Query:    https://bazaar.abuse.ch/sample/{file_info.get('sha256', '')}\n"
            f"ThreatFox Query:        https://threatfox.abuse.ch/browse/?query={file_info.get('imphash', '')}"
        )
        set_data_row(28, osint_text)

        # ----------------------------------------------------
        # 3. BEHAVIORAL ANALYSIS (Rows 31 - 37)
        # ----------------------------------------------------
        # Row 31: File System & Registry Artifacts
        host_lines = []
        dropped = host_beh.get("dropped_files", [])
        if dropped:
            host_lines.append(f"DROPPED / WRITTEN EXECUTABLE FILES ({len(dropped)}):")
            for df in dropped[:8]:
                host_lines.append(f"  [{df.get('process')}] -> {df.get('path')} ({df.get('operation')})")
            host_lines.append("")

        pers = host_beh.get("persistence_registry", [])
        if pers:
            host_lines.append(f"REGISTRY PERSISTENCE MODIFICATIONS ({len(pers)}):")
            for pr in pers[:8]:
                host_lines.append(f"  [{pr.get('process')}] {pr.get('operation')}: {pr.get('key_path')} (Detail: {pr.get('detail')})")
            host_lines.append("")

        spawned = host_beh.get("spawned_processes", [])
        if spawned:
            host_lines.append(f"SPAWNED PROCESSES ({len(spawned)}):")
            for sp in spawned[:6]:
                host_lines.append(f"  Parent: {sp.get('parent_process')} -> Command: {sp.get('command_line')}")

        if not host_lines:
            host_lines.append("No active file drop or registry modification logged in provided telemetry.")
        set_data_row(31, "\n".join(host_lines))

        # Row 33: Network Artifacts
        net_lines = []
        dns_queries = network.get("dns_queries", [])
        if dns_queries:
            net_lines.append(f"DNS RESOLUTION ACTIVITY ({len(dns_queries)} queries):")
            for d in dns_queries[:8]:
                ips_str = ", ".join(d.get("resolved_ips", [])) if d.get("resolved_ips") else "No answer"
                net_lines.append(f"  Domain: {d.get('domain')} -> Resolved IPs: [{ips_str}]")
            net_lines.append("")

        http_reqs = network.get("http_requests", [])
        if http_reqs:
            net_lines.append(f"HTTP/S C2 REQUESTS ({len(http_reqs)}):")
            for hr in http_reqs[:6]:
                net_lines.append(f"  {hr.get('method')} http://{hr.get('host')}{hr.get('uri')} (UA: {hr.get('user_agent', 'None')})")
            net_lines.append("")

        tls_snis = network.get("tls_sni", [])
        if tls_snis:
            net_lines.append(f"TLS SNI ENCRYPTED DESTINATIONS ({len(tls_snis)}):")
            for sni in tls_snis[:6]:
                net_lines.append(f"  SNI: {sni}")

        if not net_lines:
            net_lines.append("No network traffic captured in provided PCAP trace.")
        set_data_row(33, "\n".join(net_lines))

        # Row 35: Memory Analysis
        anomalies = static_data.get("anomalies", [])
        mem_lines = []
        if caps.get("Process Injection (T1055)"):
            mem_lines.append(f"Process Injection APIs detected: {', '.join(caps['Process Injection (T1055)'])}")
            mem_lines.append("Risk of unbacked executable memory allocation and remote thread execution.")
        for anom in anomalies:
            if "RWX" in anom:
                mem_lines.append(f"Memory Permission Anomaly: {anom}")
        if not mem_lines:
            mem_lines.append("No memory allocation anomalies or injection patterns identified in static/behavioral trace.")
        set_data_row(35, "\n".join(mem_lines))

        # Row 37: Open Source Research (Network)
        remote_ips = network.get("remote_ips", [])
        osint_net = []
        for r_ip in remote_ips[:5]:
            osint_net.append(f"IP WHOIS / AbuseIPDB: https://www.abuseipdb.com/check/{r_ip.get('ip')}")
        if not osint_net:
            osint_net.append("No external destination IPs available for OSINT querying.")
        set_data_row(37, "\n".join(osint_net))

        # ----------------------------------------------------
        # 4. CODE ANALYSIS (Rows 40 - 42)
        # ----------------------------------------------------
        # Row 40: Static Analysis (IDA Pro / Disassembly Highlights)
        code_lines = [
            f"Entry Point: {file_info.get('entry_point', '0x0')} (ImageBase: {file_info.get('image_base', '0x0')})",
            f"Control Flow & Architecture: {file_info.get('architecture', 'x86_64')}"
        ]
        if api_hashes:
            code_lines.append(f"API Hashing Technique (Maldev Module 55):")
            code_lines.append(f"  The binary employs custom GetProcAddress/GetModuleHandle wrappers to dynamically resolve APIs.")
            for m in api_hashes[:6]:
                code_lines.append(f"  - Hash {m['hash_hex']} -> {m['api']} using algorithm {m['algorithm']}")
        else:
            code_lines.append("Standard direct function CALLs through Import Address Table.")
        set_data_row(40, "\n".join(code_lines))

        # Row 42: Debugging (OllyDbg / x64dbg)
        dbg_lines = [
            f"OEP (Original Entry Point): {file_info.get('entry_point', '0x0')}",
            "Unpacking notes: If high section entropy is present, place execution breakpoint on VirtualAlloc/VirtualProtect and run to OEP transition."
        ]
        set_data_row(42, "\n".join(dbg_lines))

        # ----------------------------------------------------
        # 5. ANALYSIS SUMMARY (Rows 45 - 57)
        # ----------------------------------------------------
        # Row 45: Key IOCs
        iocs = ai_data.get("key_iocs", {})
        ioc_lines = ["HOST-BASED INDICATORS (IOCs):"]
        for h in iocs.get("host_iocs", [])[:10]:
            ioc_lines.append(f"  [Host] {h}")
        ioc_lines.append("\nNETWORK-BASED INDICATORS (IOCs):")
        for n in iocs.get("network_iocs", [])[:10]:
            ioc_lines.append(f"  [Net]  {n}")
        set_data_row(45, "\n".join(ioc_lines))

        # Row 47: Key Functionality
        set_data_row(47, ai_data.get("key_functionality", "Execution of sample with observed behavioral capabilities."))

        # Row 49: Purpose
        set_data_row(49, ai_data.get("purpose", "Establish foothold and execute remote commands."))

        # Row 51: Persistence
        set_data_row(51, ai_data.get("persistence", "No persistence observed."))

        # Row 53: Environment Impact
        set_data_row(53, ai_data.get("environment_impact", "Standard system compromise risk."))

        # Row 55: Root Cause
        set_data_row(55, ai_data.get("root_cause", "Phishing or malicious download."))

        # Row 57: Attribution & Recommendations
        mitre_list = ai_data.get("mitre_attack", [])
        mitre_summary = []
        if mitre_list:
            mitre_summary.append("MITRE ATT&CK MAPPING:")
            for m in mitre_list[:6]:
                mitre_summary.append(f"  {m.get('technique_id')}: {m.get('technique_name')} ({m.get('tactic')})")
            mitre_summary.append("")

        recs = ai_data.get("incident_recommendations", [])
        rec_lines = []
        if recs:
            rec_lines.append("INCIDENT RESPONSE RECOMMENDATIONS:")
            for r in recs:
                rec_lines.append(f"  - {r}")

        attrib_text = (
            f"Attribution: {ai_data.get('attribution', 'Unattributed')}\n"
            f"Threat Level: {ai_data.get('threat_level', 'UNKNOWN')} (Score: {ai_data.get('threat_score', 0)}/100)\n\n"
            + "\n".join(mitre_summary) + "\n" + "\n".join(rec_lines)
        )
        set_data_row(57, attrib_text)

        # Save document
        doc.save(str(output_path))
        return output_path
