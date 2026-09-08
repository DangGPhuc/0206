"""
Automated Pipeline Verification Tests
Tests static PE parsing, API hash detection, behavioral log ingestion,
and DOCX report generation.
"""
import unittest
from pathlib import Path
import docx

from analyzer.static import PEStaticAnalyzer
from analyzer.behavioral import BehavioralAnalyzer
from analyzer.api_hash_db import scan_binary_for_api_hashes, hash_djb2
from ai.agent import LLMThreatSynthesizer
from reporter.docx_generator import FOR610ReportGenerator

TESTS_DIR = Path(__file__).resolve().parent

class TestAutoSleuthPipeline(unittest.TestCase):

    def setUp(self):
        self.sample_pe = TESTS_DIR / "sample_benign_triage.exe"
        self.sample_pcap = TESTS_DIR / "sample_traffic.pcap"
        self.sample_procmon = TESTS_DIR / "sample_procmon.csv"
        self.output_docx = TESTS_DIR / "unit_test_report.docx"

    def test_01_api_hashing_detection(self):
        """Verify API hash database calculates and detects expected hash constants."""
        h_val = hash_djb2("VirtualAlloc")
        self.assertEqual(h_val, 0x382C0F97)

        if self.sample_pe.exists():
            with open(self.sample_pe, "rb") as f:
                data = f.read()
            matches = scan_binary_for_api_hashes(data)
            apis = [m["api"] for m in matches]
            self.assertIn("VirtualAlloc", apis)

    def test_02_static_pe_analyzer(self):
        """Verify static PE analysis extracts headers, hashes, and sections."""
        if not self.sample_pe.exists():
            self.skipTest("Sample PE not compiled yet.")
        analyzer = PEStaticAnalyzer(self.sample_pe)
        res = analyzer.analyze()
        self.assertTrue(res["file_info"]["is_pe"])
        self.assertEqual(res["file_info"]["architecture"], "64-bit (x64)")
        self.assertGreater(len(res["sections"]), 0)
        self.assertIn(".text", [s["name"] for s in res["sections"]])

    def test_03_behavioral_analyzer(self):
        """Verify network PCAP and Procmon CSV parsing."""
        if not self.sample_pcap.exists() or not self.sample_procmon.exists():
            self.skipTest("Sample PCAP or Procmon CSV missing.")
        analyzer = BehavioralAnalyzer(pcap_path=self.sample_pcap, procmon_path=self.sample_procmon)
        res = analyzer.analyze()
        
        # Check PCAP findings
        domains = [d["domain"] for d in res["network"]["dns_queries"]]
        self.assertIn("malicious-c2.test", domains)
        self.assertGreater(len(res["network"]["http_requests"]), 0)

        # Check Procmon findings
        dropped_paths = [df["path"] for df in res["host_behavior"]["dropped_files"]]
        self.assertTrue(any("dropped_payload.exe" in p for p in dropped_paths))
        self.assertGreater(len(res["host_behavior"]["persistence_registry"]), 0)

    def test_04_docx_report_generation(self):
        """Verify generic report generator and optional SANS-style adapter."""
        static_data = PEStaticAnalyzer(self.sample_pe).analyze()
        beh_data = BehavioralAnalyzer(self.sample_pcap, self.sample_procmon).analyze()
        ai_data = LLMThreatSynthesizer(provider="offline").synthesize({"static": static_data, "behavioral": beh_data})

        # Test generic built-in generator
        reporter = FOR610ReportGenerator()
        out_path = reporter.generate(static_data, beh_data, ai_data, self.output_docx)
        self.assertTrue(out_path.exists())
        doc = docx.Document(str(out_path))
        self.assertGreater(len(doc.tables), 0)

        # Test optional SANS template if present on host
        sans_tpl = Path("/run/media/kali/New Volume/malware/Malware_Analysis_Report_Template.docx")
        if sans_tpl.exists():
            sans_out = TESTS_DIR / "sans_test_report.docx"
            sans_reporter = FOR610ReportGenerator(template_path=sans_tpl)
            sans_res = sans_reporter.generate(static_data, beh_data, ai_data, sans_out)
            self.assertTrue(sans_res.exists())
            sans_doc = docx.Document(str(sans_res))
            self.assertEqual(len(sans_doc.tables[0].rows), 58)
            self.assertIn("sample_benign_triage.exe", sans_doc.tables[0].rows[3].cells[2].text)


if __name__ == "__main__":
    unittest.main()
